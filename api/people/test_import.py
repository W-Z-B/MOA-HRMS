from datetime import date
from io import StringIO

import pytest
from django.core.management import CommandError, call_command
from openpyxl import Workbook

from leave.models import LeaveLedger
from org.models import OrgUnit, Position
from people.models import Assignment, Employee


def _workbook(path, *, bad_campus=False):
    wb = Workbook()
    wb.remove(wb.active)
    units = wb.create_sheet("Units")
    units.append(["code", "name", "unit_type", "campus_code", "parent_code"])
    units.append(["AGR", "Agriculture Department", "department", "MRP", ""])
    units.append(["LIV", "Livestock Unit", "unit", "MRP", "AGR"])
    positions = wb.create_sheet("Positions")
    positions.append(
        [
            "number",
            "title",
            "unit_code",
            "scale_code",
            "grade_code",
            "step",
            "amount",
            "effective_from",
            "status",
        ]
    )
    positions.append(
        ["LIV-001", "Livestock Instructor", "LIV", "GS", "GS5", 1, 250000, "2026-01-01", "approved"]
    )
    positions.append(["LIV-002", "Farm Hand", "LIV", "GS", "GS2", 1, 120000, date(2026, 1, 1), ""])
    employees = wb.create_sheet("Employees")
    employees.append(
        [
            "employee_no",
            "first_name",
            "last_name",
            "other_names",
            "date_of_birth",
            "gender",
            "national_id",
            "nis_no",
            "tin",
            "email",
            "phone",
            "address",
            "campus_code",
            "status",
            "position_number",
            "appointment_type",
            "start_date",
            "probation_end",
        ]
    )
    employees.append(
        [
            "E0001",
            "Asha",
            "Persaud",
            "",
            "14/03/1990",
            "F",
            "123456789",
            "A1234567",
            "TIN0001",
            "asha@gsa.edu.gy",
            "600-0000",
            "Mon Repos",
            "MRP",
            "active",
            "LIV-001",
            "permanent",
            "2026-02-01",
            "",
        ]
    )
    employees.append(
        [
            "E0002",
            "Ravi",
            "Singh",
            "",
            date(1992, 6, 1),
            "M",
            "",
            "",
            "",
            "",
            "",
            "",
            "ZZZ" if bad_campus else "MRP",
            "active",
            "LIV-002",
            "seasonal",
            "2026-03-01",
            "2026-06-01",
        ]
    )
    balances = wb.create_sheet("LeaveBalances")
    balances.append(["employee_no", "leave_type_code", "days", "as_of"])
    balances.append(["E0001", "ANN", 12.5, "2026-01-01"])
    balances.append(["E0002", "ANN", 4, ""])
    wb.save(path)
    return path


@pytest.mark.django_db
def test_import_is_validated_written_and_idempotent(seeded, tmp_path):
    path = _workbook(tmp_path / "staff.xlsx")

    out = StringIO()
    call_command("import_staff", file=str(path), dry_run=True, stdout=out)
    assert "DRY RUN" in out.getvalue() and Employee.objects.count() == 0

    out = StringIO()
    call_command("import_staff", file=str(path), stdout=out)
    assert OrgUnit.objects.get(code="LIV").parent.code == "AGR"
    assert Position.objects.count() == 2
    asha = Employee.objects.get(employee_no="E0001")
    assert asha.date_of_birth == date(1990, 3, 14) and asha.nis_no == "A1234567"
    assert Assignment.objects.get(employee=asha).position.number == "LIV-001"
    assert LeaveLedger.objects.filter(reason="opening").count() == 2
    assert "A1234567" not in out.getvalue()

    out = StringIO()
    call_command("import_staff", file=str(path), stdout=out)
    text = out.getvalue()
    assert Employee.objects.count() == 2 and LeaveLedger.objects.filter(reason="opening").count() == 2
    assert "Employees" in text and "      2        0        2" in text  # read 2, created 0, updated 2


@pytest.mark.django_db
def test_import_aborts_on_validation_error_without_writing(seeded, tmp_path):
    path = _workbook(tmp_path / "bad.xlsx", bad_campus=True)
    out = StringIO()
    with pytest.raises(CommandError):
        call_command("import_staff", file=str(path), stdout=out)
    assert "Employees row 3, campus_code: unknown campus" in out.getvalue()
    assert Employee.objects.count() == 0 and OrgUnit.objects.count() == 0


@pytest.mark.django_db
def test_write_template(tmp_path):
    from openpyxl import load_workbook

    target = tmp_path / "template.xlsx"
    call_command("import_staff", write_template=str(target), stdout=StringIO())
    assert set(load_workbook(target).sheetnames) >= {
        "Units",
        "Positions",
        "Employees",
        "LeaveBalances",
        "Notes",
    }
